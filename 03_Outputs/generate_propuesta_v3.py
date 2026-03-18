#!/usr/bin/env python3
"""Generate Propuesta Triviapp V3 (Lean) (.docx)"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# -- Styles --
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for level in range(1, 4):
    h = doc.styles[f'Heading {level}']
    h.font.name = 'Calibri'
    h.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(14)
doc.styles['Heading 3'].font.size = Pt(12)

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def add_table(headers, rows, bold_last_row=False):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        is_last = (ri == len(rows) - 1) and bold_last_row
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    if is_last:
                        r.bold = True
    doc.add_paragraph()


def add_note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


# ============================================================
# TITLE
# ============================================================
title = doc.add_heading('Propuesta Triviapp V3 (Lean)', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run('Preparado por: ').bold = True
meta.add_run('Nicolás Lundin\n')
meta.add_run('Cliente: ').bold = True
meta.add_run('Marcelo "Barri" Ibáñez\n')
meta.add_run('Fecha: ').bold = True
meta.add_run('Marzo 2026')

doc.add_paragraph()

# ============================================================
# 1. CONTEXTO Y PROCESO
# ============================================================
doc.add_heading('1. Cómo llegamos hasta acá', level=1)

doc.add_heading('Propuesta v2 (Febrero 2026)', level=2)
doc.add_paragraph(
    'La propuesta inicial planteaba dos opciones de desarrollo para una plataforma '
    'completa de trivia con CMS, panel de administración, theming configurable y '
    'motor reutilizable para múltiples torneos:'
)
add_table(
    ['', 'Básica', 'Pro'],
    [
        ['Dev', '80h', '160h'],
        ['Branding + DS', '30h', '40h'],
        ['Adaptación visual por torneo', '+10h c/u', '+10h c/u'],
        ['Total 3 pilotos', '130h · $3.575.000', '220h · $6.050.000'],
        ['Para arrancar', '110h · $3.025.000', '200h · $5.500.000'],
    ],
)
add_note('Valores en CLP brutos ($27.500 CLP/h). Incluyen retención BH 15,25%.')

doc.add_heading('Reuniones de revisión (10 y 12 de Marzo)', level=2)
doc.add_paragraph(
    'El cliente revisó la propuesta en detalle y compartió un documento con ajustes,'
    'definiciones de mecánica, KPIs y funcionalidades nuevas. En dos reuniones '
    'posteriores se afinaron:'
)
items_reuniones = [
    'Mecánica del torneo: 45 minutos simultáneo, self-paced, 30 segundos por pregunta',
    'Bloques de dificultad: F-F-M-D-M con preguntas sociales cada 2 bloques',
    'Formatos de pregunta: texto, imagen, audio (clips de canción), texto libre (sociales)',
    'Ranking solo al final — sin ranking real-time',
    'Motor adaptativo eliminado — heurística estática suficiente para pilotos',
    'D1/D7 eliminado — no aplica al modelo de evento único',
    'Friends & family obligatorio pre-torneo',
    'Audio como formato de pregunta confirmado (sin impacto presupuestario)',
    'Sound design UI diferido — guiños temáticos puntuales sí incluidos',
    'Error reporting: canal externo (WhatsApp o plugin), no desarrollo in-app',
    'Anti-cheat: preguntas servidas en lotes pequeños (no todo el banco)',
    'Sesión única concurrente por usuario',
    'Requerimientos mínimos comunicados antes de compra',
    'Adaptación visual por torneo: el cliente prefiere pagar para que se haga, no gestionar vía admin',
]
for item in items_reuniones:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('El replanteo', level=2)
doc.add_paragraph(
    'Al cruzar la propuesta con las reuniones y ajustar los costos, surgió una '
    'pregunta fundamental: ¿tiene sentido construir una plataforma completa antes '
    'de saber si la gente va a comprar un ticket?'
)
doc.add_paragraph(
    'La opción Básica invierte ~$3MM en admin, CMS y theming que el cliente no va a '
    'operar directamente. La opción Pro cuesta el doble. En ambos casos, la '
    'adaptación visual de cada torneo es un cobro adicional — el admin no elimina '
    'ese costo, solo cambia quién carga las preguntas.'
)
doc.add_paragraph(
    'Esta propuesta V3 reorienta la inversión: en vez de gastar en infraestructura '
    'backend, invertimos en lo que el jugador ve y siente. El resultado es un '
    'producto que se siente profesional y validable, a un costo significativamente menor.'
)

doc.add_page_break()

# ============================================================
# 2. COMPARATIVA
# ============================================================
doc.add_heading('2. Comparativa: Básica vs Pro vs Lean', level=1)

add_note(
    'Motor de juego, mecánica, formatos de pregunta, anti-cheat, friends & family '
    'y assets visuales (IA + stock, costo cliente) son idénticos en las tres opciones.'
)

add_table(
    ['', 'Básica', 'Pro', 'Lean'],
    [
        ['Experiencia UX',
         'Funcional — animaciones estándar',
         'Inmersiva — Rive, diseño sonoro',
         'Polished — CSS animations, micro-interactions, guiños audio'],
        ['Landing',
         'Preventa + compra tickets',
         '= Básica',
         '= Básica + demo interactiva'],
        ['Admin / CMS',
         'Panel tabular (preguntas, theming básico)',
         'Panel con preview visual',
         'Manual (gestionado por el desarrollador)'],
        ['Branding',
         'Logo + sistema de diseño básico',
         'Logo + sistema de diseño completo',
         'Logo + styleguide (sin DS)'],
        ['Seguridad',
         'Sin revisión senior',
         'Revisión dev senior (completa)',
         'Revisión dev senior (reducida)*'],
        ['Torneos adicionales',
         'Admin incluido — adaptación visual cobro aparte',
         '= Básica',
         'Clone + ajuste — precio base por torneo'],
    ],
)
add_note(
    '* Lean tiene arquitectura single-tenant (un torneo a la vez), sin panel multi-usuario, '
    'schema más simple. Menor superficie de ataque, más fácil de auditar.'
)

doc.add_page_break()

# ============================================================
# 3. PROPUESTA V3 LEAN
# ============================================================
doc.add_heading('3. Propuesta V3: Lean', level=1)

doc.add_heading('Principio', level=2)
doc.add_paragraph(
    'Construir un torneo bien hecho, validar con datos reales, y escalar solo si funciona. '
    'Sin admin, sin CMS, sin infraestructura especulativa. El ahorro se reinvierte en '
    'la experiencia que el jugador ve y siente.'
)

doc.add_heading('Qué incluye', level=2)
items_incluye = [
    'Motor de juego completo: preguntas (texto, imagen, audio), scoring graduado, ranking final con reveal',
    'Bloques de dificultad (F-F-M-D-M) + preguntas sociales con clustering LLM al cierre',
    'Cascada de desempate en 3 criterios (sin necesidad de modo HQ)',
    'Landing dedicada al torneo con módulo de compra de tickets (Webpay/Stripe)',
    'Demo interactiva en landing — 5-6 preguntas que muestran el flujo real del juego',
    'Guiños de audio temáticos (riff de intro, pantalla de resultados)',
    'Animaciones CSS y micro-interactions profesionales',
    'Optimizado para móvil (web responsive, no PWA)',
    'Friends & family pre-torneo',
    'Anti-cheat (preguntas en lotes) + logs por usuario (accountability)',
    'Revisión de seguridad dev senior (reducida — menor superficie de ataque)',
    'Branding: logo, isotipo, styleguide básica interna para retomar a futuro',
]
for item in items_incluye:
    doc.add_paragraph(item, style='List Bullet')

p = doc.add_paragraph()
p.add_run('Qué se difiere (se construye post-validación):').bold = True
items_difiere = [
    'Panel de administración / CMS',
    'Theming configurable',
    'Sistema de diseño completo (DS)',
    'PWA (service workers, installable)',
    'Rive avanzado',
    'Sound design UI completo',
    'Página institucional de Soñado',
]
for item in items_difiere:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# ============================================================
# 4. MECANICA DE SCORING
# ============================================================
doc.add_heading('4. Mecánica de scoring y dificultad', level=1)

doc.add_heading('Scoring graduado', level=2)
doc.add_paragraph(
    'En vez del modelo binario (correcto/incorrecto), cada alternativa tiene un '
    'puntaje asignado según su calidad. Esto agrega profundidad estratégica, '
    'premia conocimiento parcial y hace los empates extremadamente raros.'
)

add_table(
    ['Alternativa', 'Puntos', 'Ejemplo: ¿Quién es el baterista de Metallica?'],
    [
        ['Correcta', '+5', 'Lars Ulrich'],
        ['Cercana / relacionada', '+1', 'Jason Newsted (fue bajista, no baterista)'],
        ['Neutra', '0', 'Kirk Hammett (guitarrista, más conocido)'],
        ['Trampa', '-2', 'Dave Mustaine (ex-Metallica, nunca fue baterista)'],
    ],
)

doc.add_heading('Diseño de dificultad a través del contenido', level=2)
doc.add_paragraph(
    'La dificultad no se define en el motor — se define en las preguntas. '
    'La calidad de la trampa determina la dificultad real:'
)

add_table(
    ['Dificultad', 'Tipo de trampa', 'Ejemplo'],
    [
        ['Fácil', 'Trampa obvia — claramente no pertenece', 'Justin Bieber'],
        ['Media', 'Trampa plausible — mismo universo', 'Dave Mustaine (ex-integrante)'],
        ['Difícil', 'Trampa diabólica — nombres fusionados o confundibles',
         'Kirk Ulrich, Davos Mustaine'],
    ],
)

doc.add_paragraph(
    'Cada pregunta necesita 4 alternativas con puntaje asignado. Esto es trabajo '
    'de contenido, no de desarrollo — el motor solo lee el puntaje de cada alternativa. '
    'La experiencia se juega en la calidad de las preguntas, no en la tecnología.'
)

doc.add_heading('Cascada de desempate', level=2)
doc.add_paragraph(
    'Con scoring graduado los empates son raros, pero si ocurren:'
)

add_table(
    ['#', 'Criterio', 'Qué premia'],
    [
        ['1', 'Mayor puntaje total', 'Conocimiento + velocidad'],
        ['2', 'Menos respuestas con puntos negativos', 'Criterio estratégico'],
        ['3', 'Menor tiempo promedio en respuestas correctas (+5)', 'Eficiencia'],
    ],
)

doc.add_paragraph(
    'Tres criterios, sin necesidad de un modo HQ-style de eliminación. '
    'Se resuelve en la lógica del motor, no en features adicionales. '
    'Todos los criterios se informan al jugador antes del torneo.'
)

doc.add_heading('Bonus de tiempo', level=2)
doc.add_paragraph(
    'Los segundos restantes al responder se convierten en puntos bonus, pero solo '
    'cuando la respuesta es correcta (+5). Esto evita que el bonus de velocidad '
    'distorsione el scoring graduado — responder una trampa rápido no debe premiar '
    'más que responder correctamente con calma.'
)

add_table(
    ['Respuesta', 'En 5s', 'En 15s', 'En 25s', 'En 30s'],
    [
        ['Correcta (+5)', '5 + 25 = 30 pts', '5 + 15 = 20 pts', '5 + 5 = 10 pts', '5 + 0 = 5 pts'],
        ['Cercana (+1)', '1 pt', '1 pt', '1 pt', '1 pt'],
        ['Neutra (0)', '0 pts', '0 pts', '0 pts', '0 pts'],
        ['Trampa (-2)', '-2 pts', '-2 pts', '-2 pts', '-2 pts'],
    ],
)

doc.add_paragraph(
    'El bonus genera una tensión estratégica: si estás seguro de la respuesta, '
    'respondé rápido para maximizar puntos. Si no estás seguro, tomate el tiempo '
    'para evitar la trampa. Ambas decisiones son válidas y conscientes.'
)

doc.add_heading('Omitir preguntas', level=2)
doc.add_paragraph(
    'Para el primer torneo, no se incluye botón de omitir. Si se agotan los 30 '
    'segundos sin responder, la pregunta avanza automáticamente como 0 puntos '
    '(equivalente a respuesta neutra). Esta decisión se valida en el friends & family '
    '— si resulta tedioso, se agrega la opción de skip en iteraciones posteriores.'
)
add_note(
    'Nota estratégica: permitir skip cambiaría significativamente la dinámica — '
    'el jugador cauteloso saltaría las difíciles para llegar a más fáciles. '
    'Es mejor probar sin skip primero.'
)

doc.add_heading('Banco de preguntas', level=2)
doc.add_paragraph(
    'Con scoring graduado y bonus de tiempo, los jugadores tienden a tomarse más '
    'tiempo por pregunta (evitar trampas vale más que responder rápido sin certeza). '
    'Esto reduce significativamente la cantidad de preguntas necesarias respecto '
    'al cálculo original de 702.'
)

add_table(
    ['Perfil de jugador', 'Tiempo promedio respuesta', 'Preguntas en 45 min'],
    [
        ['Cauteloso', '25-28s', '~80-90'],
        ['Balanceado', '15-20s', '~120-140'],
        ['Speed (responde en 5s)', '5-8s', '~300-340'],
    ],
)

doc.add_paragraph(
    'El cálculo considera ~3 segundos de transición entre preguntas (feedback + carga) '
    'y preguntas sociales intercaladas cada 10 preguntas (~5s extra por tipear o leer opciones). '
    'El banco necesario se calcula como 30% sobre el máximo que un jugador rápido puede alcanzar.'
)

p = doc.add_paragraph()
p.add_run('Banco recomendado: ~440 preguntas').bold = True
p.add_run(
    ' (vs 702 del cálculo original que no consideraba transiciones ni el efecto '
    'del scoring graduado). Distribución: 40% fáciles, 30% medias, 20% difíciles, '
    '10% sociales.'
)

doc.add_page_break()

# ============================================================
# 5. PRESUPUESTO
# ============================================================
doc.add_heading('5. Presupuesto', level=1)

doc.add_heading('Para arrancar (Torneo 1)', level=2)

add_table(
    ['Ítem', 'Horas', 'CLP brutos'],
    [
        ['Branding (logo, isotipo, styleguide)', '20h', '$550.000'],
        ['Desarrollo Torneo 1', '60h', '$1.650.000'],
        ['Total para arrancar', '80h', '$2.200.000'],
    ],
    bold_last_row=True,
)

doc.add_heading('3 pilotos completos', level=2)

add_table(
    ['Ítem', 'Horas', 'CLP brutos'],
    [
        ['Branding', '20h', '$550.000'],
        ['Torneo 1 (desde cero)', '60h', '$1.650.000'],
        ['Torneo 2 (clone + ajuste visual)', '14h', '$385.000'],
        ['Torneo 3 (clone + ajuste visual)', '14h', '$385.000'],
        ['Total 3 pilotos', '108h', '$2.970.000'],
    ],
    bold_last_row=True,
)

add_note(
    'Tarifa: $27.500 CLP/h bruto (incluye retención BH 15,25%). '
    'Forma de pago: 50% inicio, 50% contra entrega, por torneo. '
    'Assets visuales (fotografía stock o propia) a cargo del cliente — '
    'referencia: $1-$15 USD/imagen en plataformas de stock. '
    'Con 5-8 fotos se cubre un torneo.'
)

doc.add_paragraph(
    'Cada torneo adicional se clona y ajusta con un precio base de $385.000 CLP (14h). '
    'Si un torneo requiere mecánicas o lógicas nuevas (ej: tipo de pregunta especial), '
    'se cotiza como customización adicional sobre esa base.'
)

# ============================================================
# 6. TIMELINE
# ============================================================
doc.add_heading('6. Timeline estimado', level=1)

add_note('Estimación con margen. Sujeto a disponibilidad de contenido y assets por parte del cliente.')

add_table(
    ['Semana', 'Entregable'],
    [
        ['1', 'Branding: logo, isotipo, styleguide básica. Setup infraestructura (Supabase, deploy).'],
        ['2-3', 'Motor core: scoring graduado, bloques, preguntas sociales, timer, ranking.'],
        ['3-4', 'Landing: preventa, módulo de pago (Webpay/Stripe), demo interactiva.'],
        ['4', 'Integración: audio, anti-cheat, sesión concurrente, revisión dev senior.'],
        ['5', 'QA + Friends & Family: test controlado, correcciones, planes de contingencia.'],
        ['5+', 'Torneo 1 en vivo.'],
    ],
)

doc.add_paragraph(
    'Torneos 2 y 3 se ejecutan después del primero, con un estimado de '
    '1-2 semanas cada uno para clone, ajuste visual y carga de contenido.'
)

doc.add_page_break()

# ============================================================
# 7. KPIs DE VALIDACION
# ============================================================
doc.add_heading('7. KPIs de validación', level=1)

doc.add_paragraph(
    '¿Qué significa "validar"? Estos son los indicadores que se miden en cada '
    'piloto para determinar si el modelo funciona:'
)

doc.add_heading('Fase 1 — Preventa y registro', level=2)
items_kpi_1 = [
    'Tickets vendidos',
    'Tasa de conversión landing → compra',
    'Tasa de abandono en checkout',
    'CAC real (costo de adquisición por cliente — ecuación del cliente con inversión publicitaria)',
]
for item in items_kpi_1:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Fase 2 — Día del torneo', level=2)
items_kpi_2 = [
    'Show-up rate (inscritos que se conectan)',
    'Tasa de finalización del torneo',
    'Tiempo de respuesta promedio por jugador',
    'Número de aciertos y errores por jugador',
    '% de aciertos por nivel de dificultad',
    'Tasa de abandono durante el torneo',
    'Errores técnicos reportados',
]
for item in items_kpi_2:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Fase 3 — Post torneo', level=2)
items_kpi_3 = [
    'Tasa de intención de repetir (social / mail)',
    'Reclamos o disputas sobre el resultado',
    'Menciones orgánicas en redes post torneo',
]
for item in items_kpi_3:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Fase 4 — Rendimiento técnico', level=2)
items_kpi_4 = [
    'Tiempo de carga inicial',
    'Latencia entre pregunta y respuesta registrada',
    'Uptime durante el torneo',
    'Compatibilidad de dispositivos (iOS / Android / desktop)',
]
for item in items_kpi_4:
    doc.add_paragraph(item, style='List Bullet')

add_note(
    'Los KPIs de Fase 1 y 3 son responsabilidad del cliente (marketing, redes). '
    'Los de Fase 2 y 4 se registran automáticamente en el motor y los logs.'
)

doc.add_page_break()

# ============================================================
# 8. PREREQUISITOS
# ============================================================
doc.add_heading('8. Para arrancar necesitamos', level=1)

doc.add_paragraph(
    'Antes de iniciar el desarrollo, el cliente debe tener definido o en proceso:'
)

items_prereqs = [
    'Temática del primer torneo y comunidad target',
    'Banco de preguntas: ~440 preguntas con 4 alternativas y puntaje asignado por alternativa '
    '(distribución: 40% fáciles, 30% medias, 20% difíciles, 10% sociales)',
    'Assets fotográficos: 5-8 fotos bajo licencia o propias para el torneo',
    'Definición de premio y coordinación con operadora de viajes (si aplica)',
    'Marco legal: términos y condiciones del torneo (abogado civilista), '
    'clasificación como juego de habilidad (no azar), límites de propiedad intelectual '
    '(uso de marcas, imágenes, clips de audio)',
    'Estructura tributaria definida',
    'Pasarela de pagos: cuenta Webpay o Stripe activa',
]
for item in items_prereqs:
    doc.add_paragraph(item, style='List Bullet')

add_note(
    'El desarrollo puede arrancar en paralelo con algunos de estos puntos, '
    'pero el torneo no puede ejecutarse sin todos resueltos.'
)

doc.add_page_break()

# ============================================================
# 9. CONDICIONES GENERALES
# ============================================================
doc.add_heading('9. Condiciones generales', level=1)

items_condiciones = [
    'Forma de pago: 50% al inicio, 50% contra entrega, por torneo.',
    'Propiedad intelectual: se ceden todos los derechos patrimoniales del código, '
    'diseño y marca al cliente una vez finalizado el pago.',
    'Costos de infraestructura a cargo del cliente: ~$30 USD/mes (Supabase, dominio).',
    'Costo operativo LLM para preguntas sociales: ~$0.50 USD por torneo (clustering al cierre).',
    'Marco legal: la trivia se clasifica como juego de habilidad (no azar), '
    'lo que la diferencia legalmente de rifas y sorteos.',
    'Cláusula de flexibilidad: si durante el desarrollo surgen nuevos requerimientos '
    'que excedan la bolsa inicial, se acordará una ampliación de horas bajo la misma tarifa base.',
]
for item in items_condiciones:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# ============================================================
# 10. CAMINO A FUTURO
# ============================================================
doc.add_heading('10. Camino a futuro: de Lean a Pro', level=1)

doc.add_paragraph(
    'Si los pilotos validan el modelo y se asegura inversión, la plataforma se '
    'escala a nivel Pro. Como el upgrade se hace sobre un motor probado en producción '
    'y con datos reales de 3 torneos, se construye exactamente lo que se necesita — '
    'sin features especulativos ni rework.'
)

add_table(
    ['Fase', 'Horas', 'CLP brutos'],
    [
        ['Lean — 3 pilotos', '108h', '$2.970.000'],
        ['Upgrade a Pro (admin, CMS, theming, DS, PWA, Rive, sound design)', '85h', '$2.337.500'],
        ['Total acumulado', '193h', '$5.307.500'],
    ],
    bold_last_row=True,
)

doc.add_paragraph()

p = doc.add_paragraph()
p.add_run('Comparación: ').bold = True
p.add_run(
    'Pro desde el día 1 costaría 220h · $6.050.000. '
    'El camino Lean → Pro ahorra 27h ($742.500) porque se construye informado, '
    'no sobre suposiciones. Mismo destino, menor costo, menor riesgo.'
)

doc.add_heading('Qué incluye el upgrade', level=2)
items_upgrade = [
    'Admin/CMS: panel completo para gestión de preguntas y torneos',
    'Theming configurable (colores, tipografía, dark/light)',
    'Sistema de diseño completo (upgrade desde styleguide)',
    'PWA (service workers, installable, offline)',
    'Rive avanzado (upgrade desde CSS animations)',
    'Sound design UI completo',
    'Integración y QA',
]
for item in items_upgrade:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# ============================================================
# 11. RESUMEN EJECUTIVO
# ============================================================
doc.add_heading('11. Resumen', level=1)

add_table(
    ['', 'Básica', 'Pro', 'Lean'],
    [
        ['Para arrancar', '$3.025.000', '$5.500.000', '$2.200.000'],
        ['3 pilotos', '$3.575.000', '$6.050.000', '$2.970.000'],
        ['Con plataforma completa', '—', '$6.050.000', '$5.307.500 (upgrade)'],
        ['Experiencia jugador', 'Funcional', 'Inmersiva', 'Polished'],
        ['Admin / CMS', 'Incluido', 'Incluido', 'Manual → upgrade'],
        ['Riesgo', 'Medio', 'Alto', 'Bajo'],
    ],
)

doc.add_paragraph(
    'La opción Lean permite arrancar con $2.200.000, validar con datos reales, '
    'y solo invertir en plataforma cuando el modelo esté probado. '
    'Si todo funciona, el camino Lean → Pro cuesta menos que Pro desde el día 1.'
)

doc.add_page_break()

# ============================================================
# ANEXO: DEV BREAKDOWN
# ============================================================
doc.add_heading('Anexo: Desglose de horas — Torneo 1 (60h)', level=1)

add_note('Referencia interna de distribución de esfuerzo. Sujeto a ajustes durante el desarrollo.')

add_table(
    ['Componente', 'Horas', 'Nota'],
    [
        ['Motor core (scoring, bloques, sociales, LLM, desempate, ranking)', '25h', 'Mismo motor para las 3 opciones'],
        ['Landing + módulo compra tickets', '8h', 'Webpay/Stripe'],
        ['Demo interactiva', '6h', 'Reutiliza componentes del motor'],
        ['UI/UX (CSS animations, micro-interactions)', '5h', 'Polished, no Rive'],
        ['Guiños audio temáticos', '2h', '2-3 clips por torneo'],
        ['Infraestructura + Supabase', '5h', 'Schema simple, auth, deploy'],
        ['Anti-cheat + validación server', '2h', 'Preguntas en lotes'],
        ['Revisión dev senior (reducida)', '3h', 'Menor superficie de ataque'],
        ['QA + Friends & Family', '4h', 'Test controlado pre-torneo'],
        ['Total', '60h', ''],
    ],
    bold_last_row=True,
)

# ============================================================
# FOOTER
# ============================================================
doc.add_paragraph()
doc.add_paragraph()
p_footer = doc.add_paragraph()
p_footer.add_run(
    'Este documento es un borrador de trabajo para discusión. '
    'Los valores y alcances quedan sujetos a confirmación en reunión.'
).italic = True

# Save
output_path = '/Users/nlundin/Dev/repos/pds--triviapp/03_Outputs/Propuesta_Triviapp_V3_Lean.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
